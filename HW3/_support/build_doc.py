from pathlib import Path
import json,re
from PIL import Image
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
P=Path(__file__).parent; O=P.parent; S=O/'Screenshots'; R=json.loads((P/'results.json').read_text())
# Crops retain actual screenshot pixels and remove unrelated application chrome.
for name,box in {'P1_Construction':(0,240,670,602),'P1_Production':(0,240,1240,605),'P1_Transport':(0,240,1430,637),'P2_Schedule':(0,240,1290,870)}.items():
 im=Image.open(S/(name+'_Sheet.png'));scale=im.width/2048
 im.crop(tuple(round(v*scale) for v in box)).save(P/(name+'_crop.png'))
d=Document();sec=d.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
sec.header_distance=sec.footer_distance=Inches(.492)
# Standard business brief; named overrides: black headings, compact equations/captions.
for name,size,before,after in [('Normal',11,0,6),('Title',23,0,8),('Subtitle',11,0,6),('Heading 1',16,16,8),('Heading 2',13,12,6),('Heading 3',12,8,4),('Caption',9,3,5)]:
 st=d.styles[name];st.font.name='Calibri';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
 st.paragraph_format.space_before=Pt(before);st.paragraph_format.space_after=Pt(after);st.paragraph_format.line_spacing=1.1
 if name.startswith('Heading'):st.font.bold=True
 if name in ['Caption','Normal','Title','Subtitle']:st.font.bold=False;st.font.italic=False
 for border in st.element.findall('.//'+qn('w:pBdr')):border.getparent().remove(border)
for name,font,size in [('Equation','Cambria Math',11),('Code','Consolas',9),('Table Text','Calibri',10)]:
 st=d.styles.add_style(name,1);st.font.name=font;st.font.size=Pt(size)
 st.paragraph_format.space_before=Pt(0);st.paragraph_format.space_after=Pt(4 if name!='Table Text' else 0);st.paragraph_format.line_spacing=1
sec.header.paragraphs[0].text='95-760 | Decision Making Under Uncertainty | Homework 3'
sec.header.paragraphs[0].style='Caption'
f=sec.footer.paragraphs[0];f.alignment=2;f.add_run('Homework 3 | ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
for run in f.runs:run.font.size=Pt(9)
def p(t='',style=None):
 assert '\u2014' not in t
 return d.add_paragraph(t,style)
def h(t):
 x=d.add_paragraph(t,'Heading 1');x.paragraph_format.space_before=Pt(0)
def sub(t):d.add_paragraph(t,'Heading 2')
def eq(t):
 x=p('','Equation');last=0
 for m in re.finditer(r'([A-Za-zΣ])_(\{[^}]+\}|\([^)]*\)|[A-Za-z0-9]+(?:,[A-Za-z0-9]+)?)',t):
  x.add_run(t[last:m.start()]+m.group(1));r=x.add_run(m.group(2).strip('{}'));r.font.subscript=True;last=m.end()
 x.add_run(t[last:])
def code(t):p(t,'Code')
def page():d.add_page_break()
def table(headers,rows,widths=None):
 widths=widths or [9360//len(headers)]*len(headers);widths[-1]=9360-sum(widths[:-1])
 t=d.add_table(rows=1,cols=len(headers));t.autofit=False;t.style='Table Grid'
 pr=t._tbl.tblPr;tw=pr.find(qn('w:tblW'));tw.set(qn('w:type'),'dxa');tw.set(qn('w:w'),'9360')
 ind=OxmlElement('w:tblInd');ind.set(qn('w:w'),'120');ind.set(qn('w:type'),'dxa');pr.append(ind)
 margins=OxmlElement('w:tblCellMar')
 for k,v in [('top',80),('bottom',80),('start',120),('end',120)]:
  el=OxmlElement('w:'+k);el.set(qn('w:w'),str(v));el.set(qn('w:type'),'dxa');margins.append(el)
 pr.append(margins)
 for gc,width in zip(t._tbl.tblGrid.gridCol_lst,widths):gc.set(qn('w:w'),str(width))
 for idx,vals in enumerate([headers]+rows):
  cells=t.rows[0].cells if idx==0 else t.add_row().cells
  for c,v,width in zip(cells,vals,widths):
   c.width=Inches(width/1440);c.text=str(v);c.paragraphs[0].style='Table Text'
   c._tc.get_or_add_tcPr().find(qn('w:tcW')).set(qn('w:w'),str(width))
   if idx==0:
    c.paragraphs[0].runs[0].bold=True
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F2F4F7');c._tc.get_or_add_tcPr().append(sh)
  cells[0]._tc.getparent().get_or_add_trPr().append(OxmlElement('w:cantSplit'))
 return t
def pic(path,width=6.5):
 x=p();x.paragraph_format.space_after=Pt(4);x.add_run().add_picture(str(path),width=Inches(width))
def evidence(key,label,sheetwidth=6.5):
 h(label+' | Excel evidence')
 p('Actual Excel worksheet, Solver Parameters, and Solver Results screenshots. The formulas and settings are documented in the workbook and on the preceding solution page.','Caption')
 pic(P/(key+'_crop.png'),sheetwidth)
 x=p();x.paragraph_format.space_after=Pt(4)
 x.add_run().add_picture(str(S/(key+'_Solver.png')),width=Inches(3.05))
 x.add_run('  ')
 x.add_run().add_picture(str(S/(key+'_Result.png')),width=Inches(3.1))
 p('Method: Simplex LP. Integer optimality: 0%. Integer constraints enforced. Unconstrained variables nonnegative.','Caption')

p('Homework 3 | Solutions','Title')
p('Name(s): [enter name(s)]\nExcel uploader: [enter name]')
h('Problem 1(a)-(c) | Construction cost only')
p('Total demand is 1,000 + 3,000 + 3,000 + 1,000 + 4,000 = 12,000 cars. The facility data used in all three models are:')
table(['Plant','Fixed cost ($)','Cost/car ($)','Minimum','Capacity'],[[i+1,f'{R["F"][i]:,}',R['C'][i],f'{R["L"][i]:,}',f'{R["U"][i]:,}'] for i in range(4)],[900,2400,1860,2100,2100])
sub('(a) Integer programming formulation')
p('Let y_i = 1 if Plant i is built, and 0 otherwise, for i = 1, 2, 3, 4. Let F_i and U_i denote the fixed cost and maximum capacity in the table.')
eq('Minimize Z = Σ_i F_i y_i')
eq('= 2,100,000y_1 + 2,800,000y_2 + 2,400,000y_3 + 2,700,000y_4')
eq('Subject to: 5,000y_1 + 7,000y_2 + 8,000y_3 + 6,000y_4 ≥ 12,000')
eq('y_i ∈ {0,1} for every plant i.')
p('This first model selects sufficient installed capacity. It does not choose production quantities, so minimum production and variable costs do not enter this model.')
sub('(b)-(c) Solver result and recommendation')
p('Build Plants 1 and 3. The minimum construction cost is $4,500,000. Their combined capacity is 13,000 cars, leaving capacity for 1,000 cars beyond forecast demand. Do not build Plants 2 and 4.')
p('No single plant can meet demand. Plants 1 and 3 form the least costly feasible pair, and every three-plant combination costs more.')
code('Problem 1: Min B27; change B22:B25 (binary); B29 >= 0.')
code('C22=B22*E5; D22=B22*B5, copied down to row 25.\nB27=SUM(D22:D25); B28=SUM(C22:C25); B29=B28-D28; D28=B17.')
page();evidence('P1_Construction','Problem 1(b)',5.4)
page();h('Problem 1(d)-(f) | Construction and production')
sub('(d) Integer programming formulation')
p('Let y_i be the binary build decision defined in part (a), and let q_i be the nonnegative integer number of cars produced at Plant i. Let c_i and L_i denote production cost per car and minimum production if built.')
eq('Minimize Z = Σ_i F_i y_i + Σ_i c_i q_i')
eq('Subject to: Σ_i q_i ≥ 12,000')
eq('L_i y_i ≤ q_i ≤ U_i y_i   for i = 1, 2, 3, 4')
eq('y_i ∈ {0,1}; q_i ∈ {0,1,2,...}.')
p('The linking constraints force production to zero at an unbuilt plant and enforce both production limits at a built plant. All production costs are positive, and this optimum produces exactly 12,000 cars.')
sub('(e)-(f) Solver result and recommendation')
table(['Plant','Build?','Cars produced','Production cost'],[['1','No','0','$0'],['2','No','0','$0'],['3','Yes','8,000','$1,240,000'],['4','Yes','4,000','$680,000']],[1000,1400,3160,3800])
p('Build Plants 3 and 4. Produce 8,000 cars at Plant 3 and 4,000 at Plant 4. Construction costs are $5,100,000 and production costs are $1,920,000, for a minimum total cost of $7,020,000.')
p('Plant 3 meets its 4,000-car minimum and reaches its 8,000-car capacity. Plant 4 produces between its 2,500-car minimum and 6,000-car capacity. Total unused capacity is 2,000 cars.')
sub('Excel implementation')
code('Min B43; change B38:C41. B38:B41 binary; C38:C41 integer.\nConstraints: H38:I41 >= 0; B47 >= 0; nonnegative variables.')
code('D38=B38*D5; E38=B38*E5; F38=B38*B5; G38=C38*C5.\nH38=C38-D38; I38=E38-C38. Copy through row 41.\nB43=SUM(F38:G41); B46=SUM(C38:C41); B47=B46-D46.\nD46=B17, where B17=SUM(B16:F16)=12000.')
p('H and I are the lower-bound and capacity slacks. Requiring both to be nonnegative implements the linking constraints exactly.','Caption')
page();evidence('P1_Production','Problem 1(e)')
page();h('Problem 1(g) | Full cost model')
p('Let x_ij be the nonnegative integer number of cars shipped from Plant i to city j. Retain the build decisions y_i and production quantities q_i. Each produced car is shipped to one of the five cities.')
table(['$/car','Pittsburgh','Chicago','New York','D.C.','Boston'],[[f'Plant {i+1}']+R['T'][i] for i in range(4)]+[['Demand']+[f'{v:,}' for v in R['D']]],[1560]*6)
p('Let t_ij be the transportation cost in the table and d_j the demand of city j. The city order is Pittsburgh, Chicago, New York, Washington, D.C., Boston.')
eq('Minimize Z = Σ_i F_i y_i + Σ_i c_i q_i + Σ_i Σ_j t_ij x_ij')
eq('Subject to: Σ_i x_ij = d_j   for every city j')
eq('q_i = Σ_j x_ij   for every plant i')
eq('L_i y_i ≤ q_i ≤ U_i y_i   for every plant i')
eq('y_i ∈ {0,1}; q_i and x_ij are nonnegative integers.')
p('City constraints meet each demand exactly. The production balances ensure that a plant produces what it ships. In Excel, q_i is calculated as the row sum of shipments, so it does not need a separate changing cell.')
sub('Excel implementation for part (h)')
code('Min B66; change B57:G60. B57:B60 binary; C57:G60 integer.\nConstraints: K57:L60 >= 0; C64:G64 = 0; nonnegative variables.')
code('H57=SUM(C57:G57); I57=B57*D5; J57=B57*E5.\nK57=H57-I57; L57=J57-H57. Copy through row 60.\nC62=SUM(C57:C60); C63=B16; C64=C62-C63.\nCopy the city formulas across through column G.')
code('B67=SUMPRODUCT(B57:B60,B5:B8)\nB68=SUMPRODUCT(H57:H60,C5:C8)\nB69=SUMPRODUCT(C57:G60,B12:F15)\nB66=SUM(B67:B69)')
page();h('Problem 1(h)-(i) | Full cost solution')
p('Build Plants 2 and 3. Produce 7,000 cars at Plant 2 and 5,000 cars at Plant 3. The minimum total cost is $7,339,000.')
table(['From','Pittsburgh','Chicago','New York','D.C.','Boston'],[['Plant 1',0,0,0,0,0],['Plant 2','1,000','3,000','2,000','1,000',0],['Plant 3',0,0,'1,000',0,'4,000'],['Plant 4',0,0,0,0,0],['Total','1,000','3,000','3,000','1,000','4,000']],[1560]*6)
sub('Cost calculation')
eq('Construction = 2,800,000 + 2,400,000 = $5,200,000')
eq('Production = 7,000(165) + 5,000(155) = $1,930,000')
eq('Plant 2 transport = 1,000(15) + 3,000(10) + 2,000(12) + 1,000(10)')
eq('= $79,000')
eq('Plant 3 transport = 1,000(30) + 4,000(25) = $130,000')
eq('Total = 5,200,000 + 1,930,000 + 209,000 = $7,339,000')
p('All city demands are met. Plant 2 uses its full 7,000-car capacity. Plant 3 produces above its 4,000-car minimum and has 3,000 cars of unused capacity. Plants 1 and 4 are not built and have zero production and shipments.')
sub('Recommendation to the company')
p('Use the full cost model for the final facility decision. Plants 2 and 3 cost $100,000 more to build and $10,000 more to operate than Plants 3 and 4 at their production-only optimum, but Plant 2 has much lower shipping costs. Optimizing all costs together selects Plants 2 and 3.')
p('The three reported objective values cover different cost categories. They should not be compared as if they were three estimates of the same total cost.')
page();evidence('P1_Transport','Problem 1(h)')
page();h('Problem 2(a) | Match scheduling model')
p('Use the assignment’s hypothetical group and visibility data. Let M = {1,...,6} be the matches and D = {1,...,14} be the available days. Match numbers identify pairings; they do not require chronological order.')
table(['Match','Pairing'],[['1','Colombia vs. Germany (C-G)'],['2','Colombia vs. Philippines (C-P)'],['3','New Zealand vs. Germany (NZ-G)'],['4','New Zealand vs. Philippines (NZ-P)'],['5','Colombia vs. New Zealand (C-NZ)'],['6','Germany vs. Philippines (G-P)']],[1400,7960])
p('Define x_md = 1 if match m is scheduled on day d, and 0 otherwise. Let v_md be its visibility score. For each team t, M_t is the set of matches involving that team:')
eq('M_C = {1,2,5}; M_NZ = {3,4,5}; M_G = {1,3,6}; M_P = {2,4,6}.')
eq('Maximize V = Σ_m Σ_d v_md x_md')
eq('Σ_d x_md = 1   for every match m   [play each match exactly once]')
eq('Σ_m x_md ≤ 2   for every day d   [at most two matches per day]')
eq('Σ_(m ∈ M_t) x_md ≤ 1   for every team t and day d')
eq('Σ_m x_m,14 = 2   [exactly two matches on the last day]')
eq('x_md ∈ {0,1}   for all matches and days.')
p('The team constraints make the two day-14 matches disjoint, so every team plays once on the final day. Give those two matches the same kickoff time. Since this model selects days only, their common kickoff time is an operational requirement rather than a separate day-assignment variable.')
p('“All matches should be played on the 14 days” is interpreted as scheduling all six matches within days 1 through 14. It cannot mean that each of the 14 days must contain a match.')
page();h('Problem 2(b)-(c) | Optimal schedule')
table(['Day','Match','Visibility'],[['5','New Zealand vs. Philippines','177'],['9','Colombia vs. Germany','192'],['13','Colombia vs. New Zealand','188'],['13','Germany vs. Philippines','181'],['14','Colombia vs. Philippines','150'],['14','New Zealand vs. Germany','150']],[1100,6460,1800])
eq('Maximum visibility = 177 + 192 + 188 + 181 + 150 + 150 = 1,038.')
p('Recommend the schedule above. Both day-14 games must begin at the same time. No matches are scheduled on the other days. Every match appears once, each team plays three matches, no team plays twice on one day, and no day has more than two games.')
p('Colombia and Germany play on days 9, 13 and 14. New Zealand and the Philippines play on days 5, 13 and 14. Thus, all four teams play on consecutive days 13 and 14. This is allowed in parts (a)-(c), but motivates the added constraint in part (d).')
sub('Excel implementation')
code('Problem 2: Max B43; change B24:G37 (binary).\nB39:G39 = 1; H24:H37 <= 2; I24:L37 <= 1; H37 = 2.')
code('H24=SUM(B24:G24)\nI24=B24+C24+F24; J24=D24+E24+F24\nK24=B24+D24+G24; L24=C24+E24+G24\nCopy these daily formulas through row 37.')
code('B39=SUM(B24:B37)\nB40=SUMPRODUCT($A$24:$A$37,B24:B37)\nB41=SUMPRODUCT(B6:B19,B24:B37)\nCopy these match formulas across through column G.\nB43=SUMPRODUCT(B6:G19,B24:G37)')
p('The visibility input matrix occupies B6:G19. The binary decision matrix uses the same day and match order in B24:G37. The workbook also includes a three-day-window check for part (d).')
page();evidence('P2_Schedule','Problem 2(b)')
page();h('Problem 2(d) | At least two days of rest')
p('Add the following rolling three-day-window constraint for every team t and every starting day k = 1,...,12:')
eq('Σ_(m ∈ M_t) [x_{m,k} + x_{m,k+1} + x_{m,k+2}] ≤ 1.')
p('A team may play at most one match in any three consecutive days. Therefore, a match on day 1 excludes days 2 and 3, and the next match can occur on day 4 or later. This is a linear constraint in the original binary variables and requires no new decision variables.')
sub('Examples')
eq('For Colombia, k = 1:')
eq('(x_1,1 + x_2,1 + x_5,1) + (x_1,2 + x_2,2 + x_5,2)')
eq('+ (x_1,3 + x_2,3 + x_5,3) ≤ 1.')
p('For k = 12, the constraint covers days 12, 13 and 14. Since every team plays on day 14, no team can play on day 12 or 13 under the new rule. The original optimal schedule is therefore infeasible once this rest requirement is added.')
sub('Excel implementation')
code('B61=SUM(I24:I26), copied across through E61 and down through row 72.\nAdd B61:E72 <= 1 to the existing Problem 2 Solver model.')
p('Columns I:L contain each team’s daily match count. The 12 rows of the rest check represent windows 1-3, 2-4, ..., 12-14. Part (d) asks for the mathematical constraint, so the workbook retains the requested base optimum for parts (b)-(c).')
page();h('Problem 2 | Visibility data and source notes')
table(['Day','C-G','C-P','NZ-G','NZ-P','C-NZ','G-P'],[[i+1]+v for i,v in enumerate(R['V'])],[1000,1393,1393,1393,1393,1394,1394])
p('Source: HW3.pdf, Problem 1 on pages 1-2 and Problem 2 on pages 2-3. All homework input values were transcribed from this file. All screenshots show the actual Excel workbook or Excel Solver dialogs. Screenshot crops remove unrelated desktop content only.','Caption')
sub('Workbook organization and reproducibility')
p('HW3_Solved.xlsx contains one visible tab per problem. Problem 1 contains all three cost models. Each model has its objective, decision cells, constraints, and formulas documented next to it.')
p('Saved Solver models are stored in hidden column M on Problem 1 and hidden column N on Problem 2. Use Solver’s Load/Save command to load M20:M25 for construction only, M48:M55 for production, M82:M89 for transportation, or N5:N13 on Problem 2 for scheduling. Unhide the column if you want to inspect the saved model cells.')
p('All models use Simplex LP with binary and integer restrictions, nonnegative variables, automatic scaling, and integer optimality tolerance of 0%. The “Ignore Integer Constraints” option is off. The reported objective values were independently checked using a separate integer optimization implementation.')
assert '\u2014' not in '\n'.join(x.text for x in d.paragraphs)
d.save(O/'HW3_Solved.docx')
print('Created',O/'HW3_Solved.docx')
