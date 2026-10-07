import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L

ROOT=Path(__file__).parent.parent
R=json.loads((ROOT/'_support'/'results.json').read_text())
wb=Workbook(); wb.remove(wb.active)
navy='24476B'; blue='245BA8'; pale='EAF1F8'; peach='FCE4D6'; green='E2EFDA'

def sheet(name,title,widths):
 w=wb.create_sheet(name);w.sheet_view.showGridLines=False;w.freeze_panes='B5'
 for i,v in enumerate(widths,1):w.column_dimensions[L(i)].width=v
 w.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(widths));w['A1']=title
 w['A1'].font=Font(name='Calibri',size=17,bold=True,color='FFFFFF');w['A1'].fill=PatternFill('solid',fgColor=navy);w.row_dimensions[1].height=29
 w.merge_cells(start_row=2,start_column=1,end_row=2,end_column=len(widths));w['A2']='Mahir Nagersheth | mnagersh | HW2 | Canvas uploader: mnagersh'
 w['A2'].font=Font(name='Calibri',size=10,color='555555')
 w.sheet_properties.pageSetUpPr.fitToPage=True
 w.page_setup.orientation='landscape';w.page_setup.paperSize=w.PAPERSIZE_A4;w.page_setup.fitToWidth=1;w.page_setup.fitToHeight=0
 w.print_options.headings=True;w.print_options.horizontalCentered=True
 w.oddFooter.center.text='DMU HW2 | mnagersh | Page &P'
 return w

def row(w,r,values):
 for c,v in enumerate(values,1):
  if v is not None:
   x=w.cell(r,c,v);x.font=Font(name='Calibri',size=11,color='000000' if isinstance(v,str) and v.startswith('=') else blue if isinstance(v,(int,float)) else '222222');x.alignment=Alignment(vertical='center',wrap_text=True)
 w.row_dimensions[r].height=23

def header(w,r,values):
 row(w,r,values)
 for x in w[r][:len(values)]:x.fill=PatternFill('solid',fgColor=navy);x.font=Font(name='Calibri',bold=True,color='FFFFFF',size=10)
 w.row_dimensions[r].height=30

def note(w,r,text,last=8,height=35):
 w.merge_cells(start_row=r,start_column=1,end_row=r,end_column=last);w.cell(r,1,text);w.cell(r,1).alignment=Alignment(wrap_text=True,vertical='center');w.cell(r,1).font=Font(name='Calibri',size=11);w.row_dimensions[r].height=height

def decisions(w,ref):
 for rr in w[ref]:
  for c in rr:c.fill=PatternFill('solid',fgColor=peach);c.font=Font(name='Calibri',size=11,color=blue)

def formula_doc(w,start,items,last=8):
 header(w,start,['Formula cell(s)','Formula / explanation'])
 for k,(cell,formula) in enumerate(items,start+1):
  w.cell(k,1,cell);w.merge_cells(start_row=k,start_column=2,end_row=k,end_column=last);w.cell(k,2,formula);w.cell(k,2).data_type='s';w.cell(k,2).font=Font(name='Consolas',size=10);w.cell(k,2).alignment=Alignment(wrap_text=True,vertical='center');w.row_dimensions[k].height=29

w=sheet('Read Me','Decision Making Under Uncertainty | Homework 2',[24,23,23,23,23,23,23,23])
note(w,4,'Submission: DMU_HW2_mnagersh. PDF to Gradescope; this single workbook to Canvas. Select the PDF pages for each question. Submit manually; no files have been uploaded.',height=40)
note(w,6,'Workbook guide: P1 Lunch; P1 Sensitivity (analytical check); P1 Demand; P2; P3 Base; P3 Pittsburgh; P3 Cleveland; P3 Columbus; P4 Mentoring. Excel Solver may add its own sensitivity report.',height=40)
note(w,8,'Blue numbers are inputs or Solver decision values; peach cells are decision variables; formulas are black. Each problem includes exact cell references, constraints, objective, and method.',height=35)
note(w,10,'Sources: HW2.pdf, Problems 1-4; Submission example.pdf, formatting and documentation requirements. All passenger routes preserve each origin-destination demand.',height=35)
note(w,12,'Optimization: continuous linear programs. All reported decisions are integral. Keep P1 continuous to obtain an Excel sensitivity report. P4 is a capacitated bipartite assignment network and has an integral LP optimum.',height=40)
note(w,14,'Model verification was performed independently with SciPy/HiGHS. Excel execution and generated report status are recorded in the PDF and solver_run_log.txt.',height=35)

w=sheet('P1 Lunch','Problem 1 | School lunch planning',[27,17,17,17,17,17,17,20])
header(w,4,['Input / variable','Dish 1','Dish 2','Dish 3'])
row(w,5,['Unit cost',4,1,3]);row(w,6,['Cooking minutes',3,2,2]);row(w,7,['Shelf space',.1,.2,.15]);row(w,8,['Minimum meals',10,10,10]);row(w,10,['Meals to prepare',35,15,10]);decisions(w,'B10:D10')
row(w,12,['Total cost','=SUMPRODUCT(B5:D5,B10:D10)']);w['B12'].number_format='$#,##0.00';w['B12'].fill=PatternFill('solid',fgColor=green)
header(w,14,['Constraint','Used / LHS','Relation','Limit / RHS','Slack'])
row(w,15,['Cooking time','=SUMPRODUCT(B6:D6,B10:D10)','<=',180,'=D15-B15'])
row(w,16,['Shelf storage','=SUMPRODUCT(B7:D7,B10:D10)','<=',8,'=D16-B16'])
row(w,17,['Total demand','=SUM(B10:D10)','>=',60,'=B17-D17'])
for r,c in [(18,'B'),(19,'C'),(20,'D')]:row(w,r,[f'Minimum Dish {r-17}',f'={c}10','>=',f'={c}8',f'=B{r}-D{r}'])
header(w,23,['Solver setting','Specification'])
note(w,24,'Objective: $B$12, Min. Changing cells: $B$10:$D$10.',height=25)
note(w,25,'Constraints: $B$15:$B$16 <= $D$15:$D$16; $B$17:$B$20 >= $D$17:$D$20.',height=25)
note(w,26,'Method: Simplex LP. Make unconstrained variables non-negative. No integer constraints; request Sensitivity report after Solve.',height=30)
formula_doc(w,28,[(w.cell(r,c).coordinate,w.cell(r,c).value) for r in [12,15,16,17,18,19,20] for c in ([2] if r==12 else [2,4,5]) if w.cell(r,c).data_type=='f'])
w.print_area='A1:H49'

w=sheet('P1 Sensitivity','Problem 1 | Analytical sensitivity check',[25,16,16,16,18,18,18,20])
note(w,4,'Analytically reconstructed sensitivity values, checked against the LP basis. The native Excel report, if generated, is an additional worksheet. Ranges vary one input at a time.',height=36)
header(w,6,['Variable','Final value','Reduced cost*','Cost coefficient','Allow. increase','Allow. decrease'])
row(w,7,['Dish 1',35,0,4,1,3]);row(w,8,['Dish 2',15,0,1,1,'Infinity']);row(w,9,['Dish 3',10,.5,3,'Infinity',.5])
note(w,10,'*When the Dish 3 minimum is treated as a variable bound, its reduced cost is 0.5. With an explicit minimum row, the same marginal can be reported as that row\'s shadow price.',height=35)
header(w,12,['Constraint','Final LHS','Shadow price','RHS','Allow. increase','Allow. decrease'])
for r,v in enumerate([['Time',155,0,180,'Infinity',25],['Storage',8,-30,8,2.5,.5],['Demand',60,7,60,5,12.5],['Dish 1 minimum',35,0,10,25,'Infinity'],['Dish 2 minimum',15,0,10,5,'Infinity'],['Dish 3 minimum',10,.5,10,10,10]],13):row(w,r,v)
note(w,21,'Binding: storage, demand, and Dish 3 minimum. A binding constraint holds with equality at the optimum. Time has 25 minutes of slack; Dish 1 and Dish 2 minima have slack 25 and 5.',height=40)
note(w,23,'Storage: one additional shelf lowers minimum cost by $30 for RHS in [7.5,10.5]. Demand: one additional required meal raises minimum cost by $7 for RHS in [47.5,65].',height=40)
note(w,25,'Dish 3 minimum: one additional required Dish 3 meal raises minimum cost by $0.50 for its RHS in [0,20]. Other parameters stay fixed in each interpretation.',height=35)
w.print_area='A1:H25'

w=sheet('P1 Demand','Problem 1(f) | Cost versus required lunches',[20,19,18,18,18,18,20,22])
header(w,4,['Breakpoint','Value','Explanation'])
row(w,5,['Minimum production',"=SUM('P1 Lunch'!B8:D8)"])
row(w,6,['Storage kink',"=('P1 Lunch'!D16-'P1 Lunch'!B7*'P1 Lunch'!B8-'P1 Lunch'!D7*'P1 Lunch'!D8)/'P1 Lunch'!C7+'P1 Lunch'!B8+'P1 Lunch'!D8"])
row(w,7,['Maximum lunches',"=('P1 Lunch'!D16-('P1 Lunch'!C7-'P1 Lunch'!B7)*'P1 Lunch'!C8-('P1 Lunch'!D7-'P1 Lunch'!B7)*'P1 Lunch'!D8)/'P1 Lunch'!B7"])
note(w,9,'Continuous LP curve for the stated parameters: cost = 80 for D <= 30; D + 50 for 30 <= D <= 47.5; 7D - 235 for 47.5 <= D <= 65. D > 65 is infeasible.',height=37)
note(w,10,'Chart formulas encode the stated basis regimes. If capacities or costs change, re-solve the LP before using these regimes. At D=65: (45,10,10), cost $220, cooking time 175.',height=37)
header(w,12,['Required lunches','Dish 1','Dish 2','Dish 3','Minimum cost','Cooking minutes','Shelves used','Status'])
for rr,dd in enumerate([i/2 for i in range(131)]+[66],13):
 row(w,rr,[dd,f'=IF(A{rr}>$B$7,"",IF(A{rr}<=$B$6,10,2*A{rr}-85))',f'=IF(A{rr}>$B$7,"",IF(A{rr}<=$B$5,10,IF(A{rr}<=$B$6,A{rr}-20,75-A{rr})))',f'=IF(A{rr}>$B$7,"",10)',f'=IF(A{rr}>$B$7,"",SUMPRODUCT(B{rr}:D{rr},\'P1 Lunch\'!$B$5:$D$5))',f'=IF(A{rr}>$B$7,"",SUMPRODUCT(B{rr}:D{rr},\'P1 Lunch\'!$B$6:$D$6))',f'=IF(A{rr}>$B$7,"",SUMPRODUCT(B{rr}:D{rr},\'P1 Lunch\'!$B$7:$D$7))',f'=IF(A{rr}>$B$7,"Infeasible","Feasible")'])
 w.cell(rr,5).number_format='$0.00'
chart=ScatterChart();chart.title='Minimum cost vs. lunch demand';chart.x_axis.title='Required lunches';chart.y_axis.title='Minimum cost ($)';chart.style=13
chart.series.append(Series(Reference(w,min_col=5,min_row=13,max_row=143),Reference(w,min_col=1,min_row=13,max_row=143),title='Continuous LP'));chart.width=23;chart.height=12
w.add_chart(chart,'J4');w.print_area='A1:H38'

w=sheet('P2','Problem 2 | Sensitivity statements',[16,16,20,22,22,22,22,22])
header(w,4,['Part','Answer','Justification'])
answers=[('a','True','For an optimal continuous LP, positive slack implies a zero dual multiplier by complementary slackness.'),('b','False','Counterexample: minimize x1 subject to x1 >= 1 and 0 <= x2 <= 1. Every (1,x2) is optimal, yet only x2 has a zero allowable coefficient change in at least one direction.'),('c','False','The joint RHS change is not guaranteed by separate ranges: 1/2 + 1/1 = 150%, exceeding the 100% rule. The objective might change by $8, but this report alone cannot establish it.'),('d','False','Resource 1 RHS can rise only to 12 under the stated shadow price; 14 exceeds that range. A $20 increase cannot be inferred.'),('e','False','The $3 marginal applies only while Resource 2 RHS is in [3,9], holding other data fixed. It is not valid for every additional unit indefinitely.'),('f','True','Final LHS equals RHS for both resources: 10=10 and 8=8; therefore both are binding.')]
for rr,(p,a,j) in enumerate(answers,5):
 row(w,rr,[p,a,j]);w.merge_cells(start_row=rr,start_column=3,end_row=rr,end_column=8);w.row_dimensions[rr].height=58
header(w,13,['Resource','Final LHS','Shadow price','RHS','Allow. increase','Allow. decrease','Unit cost'])
row(w,14,['Resource 1',10,5,10,2,4,5]);row(w,15,['Resource 2',8,3,8,1,5,4])
row(w,17,['Joint-change ratio','=1/E14+1/E15']);w['B17'].number_format='0%'
note(w,19,'The stated resource purchase costs do not extend sensitivity ranges. If this is profit before new resource purchases, an extra unit yields net gain $0 for Resource 1 and -$1 for Resource 2, within valid one-at-a-time ranges.',height=45)

origins=['Pittsburgh','Cleveland','Columbus'];dest=['Miami','Orlando','New York'];hubs=['Chicago','Atlanta']
demand=[[100,60,40],[50,50,50],[0,40,60]]
c1=[[50,100],[60,90],[70,80]]; c2=[[100,90,80],[80,70,100]]
caps1=[[150,100],[100,100],[100,80]];caps2=[[150,120,100],[120,150,120]]
cd=[[120,80,70],[110,75,65],[100,70,60]];ud=[[100,60,40],[100,60,50],[80,40,60]]
for scenario in ['Base']+origins:
 w=sheet('P3 '+scenario,'Problem 3 | '+scenario+' flight network',[23,19,17,16,16,16,18,17,17,18,17,18])
 note(w,3,'Direct flights allowed from: '+('none' if scenario=='Base' else scenario)+'. Each row preserves one origin-destination demand.',last=12,height=25)
 header(w,5,['Origin','Destination','Demand','Via Chicago','Via Atlanta','Direct','Assigned','Residual','Cost Chicago','Cost Atlanta','Cost direct','Row cost'])
 for i,o in enumerate(origins):
  for j,d in enumerate(dest):
   r=6+3*i+j; vals=R['p3'][scenario]['paths'][i][j]
   first=19+2*i
   row(w,r,[o,d,demand[i][j],*vals,f'=SUM(D{r}:F{r})',f'=G{r}-C{r}',f'=$D${first}+$D${25+j}',f'=$D${first+1}+$D${28+j}',cd[i][j],f'=SUMPRODUCT(D{r}:F{r},I{r}:K{r})'])
 decisions(w,'D6:F14');row(w,16,['Total cost','=SUM(L6:L14)']);w['B16'].number_format='$#,##0';w['B16'].fill=PatternFill('solid',fgColor=green)
 header(w,18,['Arc origin','Arc destination','Capacity','Unit cost','Flow','Slack'])
 for i,o in enumerate(origins):
  for h,hub in enumerate(hubs):
   r=19+2*i+h;col='D' if h==0 else 'E';start=6+3*i
   row(w,r,[o,hub,caps1[i][h],c1[i][h],f'=SUM({col}{start}:{col}{start+2})',f'=C{r}-E{r}'])
 for h,hub in enumerate(hubs):
  for j,d in enumerate(dest):
   r=25+3*h+j;col='D' if h==0 else 'E'
   row(w,r,[hub,d,caps2[h][j],c2[h][j],f'={col}{6+j}+{col}{9+j}+{col}{12+j}',f'=C{r}-E{r}'])
 header(w,32,['Direct origin','Destination','Allowed capacity','Flow','Slack'])
 for i,o in enumerate(origins):
  for j,d in enumerate(dest):
   r=33+3*i+j
   row(w,r,[o,d,ud[i][j] if scenario==o else 0,f'=F{6+3*i+j}',f'=C{r}-D{r}'])
 header(w,43,['Solver setting','Specification'])
 note(w,44,'Objective $B$16, Min; changing cells $D$6:$F$14; Simplex LP; non-negative variables.',last=12,height=25)
 note(w,45,'Constraints: $G$6:$G$14 = $C$6:$C$14; $E$19:$E$30 <= $C$19:$C$30; $D$33:$D$41 <= $C$33:$C$41.',last=12,height=27)
 note(w,46,'Connecting path costs = first-leg cost + second-leg cost. Flow formulas aggregate all OD passengers using each flight. Capacities and OD demands stay fixed across scenarios.',last=12,height=30)
 formula_doc(w,48,[('G6 (fill to G14)','=SUM(D6:F6)'),('H6 (fill to H14)','=G6-C6'),('I6, J6','=$D$19+$D$25 ; =$D$20+$D$28'),('L6 (fill to L14)','=SUMPRODUCT(D6:F6,I6:K6)'),('B16','=SUM(L6:L14)'),('E19, E20','=SUM(D6:D8) ; =SUM(E6:E8)'),('E25','=D6+D9+D12'),('F19 (fill to F30)','=C19-E19'),('D33, E33','=F6 ; =C33-D33')],last=12)
 w.print_area='A1:L58'

w=sheet('P4 Mentoring','Problem 4 | Mentor assignment',[25,16,16,16,16,16,16,17,17,17])
names=['Alice','Ben','Carlos','Diana','Emma','Frank'];score=R['p4']['scores'];chosen={(a,b) for a,b,v in R['p4']['assignments']}
header(w,4,['Score: mentor / mentee']+names)
for i,n in enumerate(names,5):row(w,i,[n]+score[i-5])
note(w,12,'A score of 0 below represents a forbidden pair, not an available zero-score match. Allowed-pair mask B26:G31 enforces the restrictions.',last=10,height=30)
header(w,14,['Assign: mentor / mentee']+names+['Mentee load','Capacity'])
for i,n in enumerate(names,15):row(w,i,[n]+[int((n,j) in chosen) for j in names]+[f'=SUM(B{i}:G{i})',3])
decisions(w,'B15:G20')
row(w,22,['Mentors assigned']+[f'=SUM({L(j)}15:{L(j)}20)' for j in range(2,8)])
row(w,23,['Required']+[1]*6)
header(w,25,['Allowed: mentor / mentee']+names)
for i,n in enumerate(names,26):row(w,i,[n]+[int(s>0) for s in score[i-26]])
row(w,33,['Total preference','=SUMPRODUCT(B5:G10,B15:G20)']);w['B33'].fill=PatternFill('solid',fgColor=green)
header(w,35,['Solver setting','Specification'])
note(w,36,'Objective $B$33, Max; changing cells $B$15:$G$20; Simplex LP; non-negative variables.',last=10,height=25)
note(w,37,'Constraints: $H$15:$H$20 <= $I$15:$I$20; $B$22:$G$22 = $B$23:$G$23; $B$15:$G$20 <= $B$26:$G$31.',last=10,height=30)
note(w,38,'This network LP has an integral optimum. Binary constraints may alternatively be added to B15:G20; the optimum remains 435.',last=10,height=26)
formula_doc(w,40,[('H15 (fill to H20)','=SUM(B15:G15)'),('B22 (fill across to G22)','=SUM(B15:B20)'),('B33','=SUMPRODUCT(B5:G10,B15:G20)')],last=10)
note(w,45,'Result: Alice mentors Diana, Emma, Frank; Ben mentors Alice and Carlos; Carlos mentors Ben. Total preference 435. The Ben-Carlos reciprocal assignment and Carlos\'s score 20 motivate fairness and cycle restrictions.',last=10,height=45)
w.print_area='A1:J45'

for w in wb:
 for rr in w:
  for c in rr:
   if c.data_type=='f':c.font=Font(name='Calibri',size=11,color='000000');c.comment=Comment('Formula documented in this worksheet. Inputs: HW2.pdf.','Mahir Nagersheth')
 w.sheet_properties.outlinePr.summaryRight=False
wb.active=0
wb.save(ROOT/'DMU_HW2_mnagersh.xlsx')
print('Saved',ROOT/'DMU_HW2_mnagersh.xlsx')
