from pathlib import Path
import numpy as np
from scipy.optimize import milp, Bounds, LinearConstraint
import openpyxl
from zipfile import ZipFile
root=Path('/Users/mahir/Downloads/DMU')
s=np.array([4,2,.5,1,1,4,9,1,.5,.5]); m=np.array([4,6,20,10,12,3,2,4,6,4]); e=np.array([10,5,.25,1.5,1,8,12,1,.5,.6]); U=np.full(10,100)
A=[]; lo=[]; hi=[]
def add(y,x,l=-np.inf,h=np.inf): A.append(np.r_[y,x]);lo.append(l);hi.append(h)
zero=np.zeros(10); I=np.eye(10)
add(s,zero,h=100);add(zero,np.ones(10),l=4)
for i in range(10): add(I[i],-m[i]*I[i],l=0);add(I[i],-U[i]*I[i],h=0)

add(zero,I[7]+I[8]+I[9]-2*I[0],l=0)
add(zero,I[0]+I[1]+I[6],l=2);add(zero,I[0]+I[6],h=1)
for i in [7,8,9]: add(zero,I[:7].sum(axis=0)-I[i],l=0)
add(I[0],-100*I[3],h=30)
r=milp(-np.r_[e,zero],integrality=np.ones(20),bounds=Bounds(np.zeros(20),np.r_[U,np.ones(10)]),constraints=LinearConstraint(A,lo,hi),options={'mip_rel_gap':0})
assert r.success and abs(-r.fun-241.6)<1e-8
w=openpyxl.load_workbook(root/'_support/w5/Recitation_W5_Solved.xlsx',data_only=True);ws=w.active
x=np.array([ws.cell(i,3).value for i in range(21,31)])
y=np.array([ws.cell(i,4).value for i in range(21,31)])
lhs=np.array(A)@np.r_[y,x]
assert np.all(lhs>=np.array(lo)-1e-7) and np.all(lhs<=np.array(hi)+1e-7)
assert np.all(np.r_[y,x]==np.round(np.r_[y,x]))
assert abs(ws['D33'].value-241.6)<1e-8 and abs(e@y-241.6)<1e-8
assert not [(c.coordinate,c.value) for row in ws for c in row if c.data_type=='e']
with ZipFile(root/'Recitation_W5_Solved.docx') as z:
 assert not any('\u2014' in z.read(n).decode('utf-8') for n in z.namelist() if n.endswith('.xml'))
 assert len([n for n in z.namelist() if n.startswith('word/media/')])==5
print('PASS: Independent MILP optimum 241.6; Excel cached solution feasible and optimal; no formula errors; 5 real screenshots embedded; no em dashes.')
