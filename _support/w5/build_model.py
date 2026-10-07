from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

root=Path('/Users/mahir/Downloads/DMU')
w=openpyxl.load_workbook(root/'Recitation W5/Recitation W5.xlsx')
s=w['Harvest']
for merged in list(s.merged_cells.ranges): s.unmerge_cells(str(merged))
s['B1']='W5 | Community garden optimization'
headers=['Food','Category','Space / plant','Minimum','Harvest / plant','Max plants']
for c,v in enumerate(headers,2): s.cell(4,c,v)
for r in range(5,15): s.cell(r,7,f'=INT($C$16/D{r})')
for c,v in enumerate(['Food','Plant (binary)','Plants (integer)','Min if planted','Max if planted','Space used','Harvest (lb)'],2): s.cell(20,c,v)
for r in range(21,31):
 p=r-16
 s.cell(r,3,0); s.cell(r,4,0)
 s.cell(r,5,f'=E{p}*C{r}')
 s.cell(r,6,f'=G{p}*C{r}')
 s.cell(r,7,f'=D{p}*D{r}')
 s.cell(r,8,f'=F{p}*D{r}')
s['B33']='Maximize expected harvest (lb)'; s['D33']='=SUM(H21:H30)'
for c,v in enumerate(['Constraint','LHS','Relation','RHS','Slack'],2): s.cell(36,c,v)
rows=[
 ('Planting space','=SUM(G21:G30)','<=','=C16','=E37-C37'),
 ('Different food types','=SUM(C21:C30)','>=',4,'=C38-E38'),
 ('Tomato requested','=C21','=',1,'=C39-E39'),
 ('Tomato companion herbs','=SUM(C28:C30)','>=','=2*C21','=C40-E40'),
 ('At least two named foods','=C21+C22+C27','>=',2,'=C41-E41'),
 ('Tomato / zucchini exclusion','=C21+C27','<=',1,'=E42-C42'),
 ('Basil needs fruit / vegetable','=SUM(C21:C27)','>=','=C28','=C43-E43'),
 ('Cilantro needs fruit / vegetable','=SUM(C21:C27)','>=','=C29','=C44-E44'),
 ('Parsley needs fruit / vegetable','=SUM(C21:C27)','>=','=C30','=C45-E45'),
 ('Over 30 tomatoes needs lettuce','=D21','<=','=30+$G$5*C24','=E46-C46'),
 ('Lettuce minimum if selected','=D24','>=','=10*C24','=C47-E47')]
for r,row in enumerate(rows,37):
 for c,v in enumerate(row,2): s.cell(r,c,v)
s['B49']='Solver: Max D33; change C21:D30; Simplex LP; integer optimality 0%.'
s['B50']='Linking: D21:D30 >= E21:E30 and D21:D30 <= F21:F30.'
s['B51']='C21:C30 binary; D21:D30 nonnegative integers.'
s['B52']='Tomatoes <= 25 from space, so the over-30 condition never activates.'
widths={'A':2,'B':32,'C':19,'D':19,'E':19,'F':19,'G':17,'H':19}
for c,x in widths.items(): s.column_dimensions[c].width=x
for row in s:
 for cell in row:
  cell.font=Font(name='Calibri',size=11,color='202B33')
  cell.alignment=Alignment(vertical='center')
  if cell.data_type=='f': cell.font=Font(name='Calibri',size=11,color='000000')
for r in range(1,53): s.row_dimensions[r].height=21
for r in [1,2,18,32,35]:
 s.merge_cells(start_row=r,start_column=2,end_row=r,end_column=8)
 s.cell(r,2).font=Font(name='Calibri',size=13,bold=True,color='FFFFFF')
 for cell in s[r][1:8]: cell.fill=PatternFill('solid',fgColor='24465B')
for r,last in [(4,7),(20,8),(36,6)]:
 for row in s.iter_rows(min_row=r,max_row=r,min_col=2,max_col=last):
  for cell in row:
   cell.fill=PatternFill('solid',fgColor='E4EDF2'); cell.font=Font(name='Calibri',size=11,bold=True)
for row in s.iter_rows(min_row=21,max_row=30,min_col=3,max_col=4):
 for cell in row: cell.fill=PatternFill('solid',fgColor='EAF2FF'); cell.font=Font(name='Calibri',size=11,color='0000FF')
s['D33'].fill=PatternFill('solid',fgColor='E2F0D9'); s['D33'].font=Font(name='Calibri',size=12,bold=True)
for row in s.iter_rows(min_row=5,max_row=47,min_col=3,max_col=8):
 for cell in row:
  if cell.value is not None and (cell.data_type=='f' or isinstance(cell.value,(int,float))): cell.number_format='0.##'
s.sheet_view.zoomScale=85
s.print_options.horizontalCentered=True
s.print_area='B1:H52'
s.page_setup.orientation='landscape'
s.page_setup.paperSize=s.PAPERSIZE_A4
s.page_setup.fitToWidth=1; s.page_setup.fitToHeight=0
w.save(root/'_support/w5/Recitation_W5_Solved.xlsx')
print('Prepared formula-driven Excel model')
